import shutil

import planilha
from util import IE_EMPRESA, IE_OUTRA


def test_ie_pb():
    assert planilha.ie_pb_valida(IE_EMPRESA)
    assert not planilha.ie_pb_valida("160123454")
    assert planilha.normalizar_ie("16.012.345-3") == (IE_EMPRESA, None)
    assert planilha.normalizar_ie(160123453.0) == (IE_EMPRESA, None)
    assert planilha.normalizar_ie("123")[0] is None


def test_cnpj_opcional():
    assert planilha.normalizar_documento("") == ("", None)
    assert planilha.normalizar_documento("11.222.333/0001-81") == ("11222333000181", None)
    assert planilha.normalizar_documento("11.222.333/0001-82")[1]


def test_planilha_modelo_do_escritorio(tmp_path):
    """Usa o arquivo modelo de verdade, preenchido como o cliente faria."""
    from openpyxl import load_workbook

    arq = tmp_path / "empresas.xlsx"
    shutil.copyfile(planilha.MODELO, arq)
    wb = load_workbook(arq)
    ws = wb.worksheets[0]
    linhas = [
        ["S", "015", "1", "11222333000181", "16.012.345-3", "Padaria Centro", "mar00134", "s3nh@", "S", "S", "N"],
        ["", "016", "", "", 160987652, "Mercado", "mar00134", "", "x", "", "sim"],
        ["N", "017", "", "", "161112226", "Inativa", "mar00134", "a", "S", "", ""],
        ["S", "018", "", "", "161112226", "Sem serviço", "mar00134", "a", "", "", ""],
        ["S", "019", "", "", "999", "IE errada", "mar00134", "a", "S", "", ""],
        ["S", "020", "", "", IE_OUTRA, "Repetida", "outro", "b", "S", "", ""],
        ["S", "021", "", "", "161112226", "Sem login", "", "b", "S", "", ""],
    ]
    for i, valores in enumerate(linhas, start=2):
        for j, v in enumerate(valores, start=1):
            ws.cell(i, j, v)
        if i > 2:  # só a primeira empresa fica com o X que o modelo traz nas outras colunas
            for j in range(len(valores) + 1, ws.max_column + 1):
                ws.cell(i, j).value = None
    wb.save(arq)

    leitura = planilha.ler(arq)
    assert [(e.codigo, e.ie, e.nome, e.login, e.tipos) for e in leitura.empresas] == [
        ("015", IE_EMPRESA, "Padaria Centro", "mar00134",
         ["nfe_saida", "nfe_entrada", "cte", "faturas", "extrato", "omissao", "certidao", "cartao", "mei"]),
        ("016", IE_OUTRA, "Mercado", "mar00134", ["nfe_saida", "nfce"]),
    ]
    assert leitura.empresas[0].cnpj == "11222333000181" and leitura.empresas[0].filial == "1"
    assert len(leitura.erros) == 2       # IE errada, sem login
    assert len(leitura.avisos) == 3      # inativa, sem serviço, repetida


def test_csv(tmp_path):
    arq = tmp_path / "e.csv"
    arq.write_text("ATIVA;ID;CNPJ;INSCRICAO ESTADUAL;NOME FANTASIA;LOGIN;SENHA;NF-e Saida\n"
                   "S;1;;160123453;Loja;mar;abc;S\n", encoding="cp1252")
    e = planilha.ler(arq).empresas[0]
    assert (e.ie, e.senha, e.tipos) == (IE_EMPRESA, "abc", ["nfe_saida"])


def test_cabecalho_errado(tmp_path):
    arq = tmp_path / "e.csv"
    arq.write_text("CNPJ;Apelido\n1;2\n")
    assert "Cabeçalho" in planilha.ler(arq).erros[0]


def test_modelo_vem_com_x_e_linhas_livres_sao_ignoradas():
    from openpyxl import load_workbook

    ws = load_workbook(planilha.MODELO).worksheets[0]
    cab = [c.value for c in ws[1]]
    assert ws.cell(2, cab.index("NF-e Saida") + 1).value == "X" and ws.cell(2, 1).value == "X"
    leitura = planilha.ler(planilha.MODELO)
    assert (leitura.empresas, leitura.erros, leitura.avisos) == ([], [], [])


def test_corrigir_nome():
    from planilha import corrigir_nome
    assert corrigir_nome("lidi@ne silv4") == "Lidiane Silva"
    assert corrigir_nome("  JOAO  DA SILVA ltda ") == "Joao da Silva LTDA"
    assert corrigir_nome("posto 2000 me") == "Posto 2000 ME"
    assert corrigir_nome("3M do brasil") == "3M do Brasil"
