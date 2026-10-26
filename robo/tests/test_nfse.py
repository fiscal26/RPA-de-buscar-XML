"""NFS-e pela API do Portal Nacional: leitura do XML, retenções, lotes por NSU, cache e relatório."""
import base64
import gzip
import json
from datetime import datetime, timedelta, timezone

import pytest

import nfse

DONO = "11222333000181"
CLIENTE = "44555666000199"
PRESTADOR = "77888999000155"
NS = "http://www.sped.fazenda.gov.br/nfse"


def nota_xml(chave, numero, prest, toma, toma_nome, valor, compet="2026-09", tp_ret_iss="1", iss=0.0,
             irrf=0.0, csll=0.0, inss=0.0, pis=0.0, cofins=0.0, tp_ret_pc="2", total_ret=None, liquido=None) -> bytes:
    total = round((iss if tp_ret_iss in ("2", "3") else 0) + irrf + csll + inss
                  + (pis if tp_ret_pc in ("1", "3") else 0) + (cofins if tp_ret_pc in ("1", "3") else 0), 2)
    total_ret = total if total_ret is None else total_ret
    liquido = round(valor - total_ret, 2) if liquido is None else liquido
    return f"""<?xml version="1.0" encoding="UTF-8"?>
<NFSe xmlns="{NS}" versao="1.00"><infNFSe Id="NFS{chave}">
 <xLocIncid>João Pessoa</xLocIncid><nNFSe>{numero}</nNFSe><cStat>100</cStat>
 <emit><CNPJ>{prest}</CNPJ><xNome>Prestadora {prest[:4]} Ltda</xNome></emit>
 <valores><vBC>{valor:.2f}</vBC><pAliqAplic>5.00</pAliqAplic><vISSQN>{iss:.2f}</vISSQN>
  <vTotalRet>{total_ret:.2f}</vTotalRet><vLiq>{liquido:.2f}</vLiq></valores>
 <DPS versao="1.00"><infDPS Id="DPS1">
  <dhEmi>{compet}-15T10:00:00-03:00</dhEmi><dCompet>{compet}-15</dCompet>
  <prest><CNPJ>{prest}</CNPJ></prest>
  <toma><CNPJ>{toma}</CNPJ><xNome>{toma_nome}</xNome></toma>
  <serv><cServ><cTribNac>010101</cTribNac><xDescServ>Consultoria</xDescServ></cServ></serv>
  <valores><vServPrest><vServ>{valor:.2f}</vServ></vServPrest>
   <trib><tribMun><tribISSQN>1</tribISSQN><pAliq>5.00</pAliq><tpRetISSQN>{tp_ret_iss}</tpRetISSQN></tribMun>
    <tribFed><piscofins><CST>01</CST><vBCPisCofins>{valor:.2f}</vBCPisCofins><vPis>{pis:.2f}</vPis>
     <vCofins>{cofins:.2f}</vCofins><tpRetPisCofins>{tp_ret_pc}</tpRetPisCofins></piscofins>
     <vRetCP>{inss:.2f}</vRetCP><vRetIRRF>{irrf:.2f}</vRetIRRF><vRetCSLL>{csll:.2f}</vRetCSLL></tribFed>
   </trib></valores>
 </infDPS></DPS>
</infNFSe></NFSe>""".encode()


def evento_xml(chave, codigo="e101101") -> bytes:
    return f"""<evento xmlns="{NS}"><infEvento Id="EVT1"><pedRegEvento><infPedReg Id="PRE1">
<chNFSe>{chave}</chNFSe><{codigo}><xDesc>Cancelamento de NFS-e</xDesc></{codigo}></infPedReg></pedRegEvento>
</infEvento></evento>""".encode()


def empacotar(xml: bytes) -> str:
    return base64.b64encode(gzip.compress(xml)).decode()


C1, C2, C3, C4, C5 = (f"2507507{i:043d}" for i in range(1, 6))


def documentos_exemplo():
    return [
        # emitida sem retenção
        (1, "NFSE", C1, nota_xml(C1, "1", DONO, CLIENTE, "Cliente Um SA", 1000.0, iss=50.0)),
        # emitida com ISS retido pelo tomador e PIS/COFINS/CSLL/IRRF retidos
        (2, "NFSE", C2, nota_xml(C2, "2", DONO, CLIENTE, "Cliente Um SA", 10000.0, tp_ret_iss="2", iss=500.0,
                                 irrf=150.0, csll=100.0, pis=65.0, cofins=300.0, tp_ret_pc="3")),
        # recebida com INSS retido
        (3, "NFSE", C3, nota_xml(C3, "77", PRESTADOR, DONO, "Lotivisse", 2000.0, inss=220.0)),
        # emitida e cancelada
        (4, "NFSE", C4, nota_xml(C4, "3", DONO, CLIENTE, "Cliente Um SA", 999.0)),
        (5, "EVENTO", C4, evento_xml(C4)),
        # outra competência
        (6, "NFSE", C5, nota_xml(C5, "4", DONO, CLIENTE, "Cliente Um SA", 500.0, compet="2026-10")),
    ]


class Resposta:
    def __init__(self, status=200, dados=None, conteudo=b"", cabecalhos=None):
        self.status_code = status
        self._dados = dados
        self.content = conteudo
        self.text = json.dumps(dados) if dados is not None else conteudo.decode(errors="ignore")
        self.headers = cabecalhos or {}

    def json(self):
        if self._dados is None:
            raise ValueError("sem JSON")
        return self._dados


class PortalFalso:
    """Imita o ADN: devolve lotes de até 2 documentos com NSU maior que o pedido."""

    def __init__(self, docs, falhas=None):
        self.docs = docs
        self.falhas = list(falhas or [])
        self.pedidos = []

    def get(self, url, params=None, headers=None, timeout=None):
        self.pedidos.append((url, dict(params or {})))
        if self.falhas:
            return self.falhas.pop(0)
        if "/danfse/" in url:
            return Resposta(200, conteudo=b"%PDF-1.4 falso")
        nsu = int(url.rsplit("/", 1)[-1])
        lote = [d for d in self.docs if d[0] > nsu][:2]
        if not lote:
            return Resposta(200, {"StatusProcessamento": "NENHUM_DOCUMENTO_LOCALIZADO", "LoteDFe": []})
        return Resposta(200, {"StatusProcessamento": "DOCUMENTOS_LOCALIZADOS", "LoteDFe": [
            {"NSU": n, "TipoDocumento": t, "ChaveAcesso": c, "ArquivoXml": empacotar(x)} for n, t, c, x in lote]})


def cliente(portal):
    return nfse.ClienteADN(portal, intervalo=0, dormir=lambda s: None)


def test_le_a_nota_e_as_retencoes():
    n = nfse.ler_nfse(documentos_exemplo()[1][3])
    assert (n.chave, n.numero, n.competencia, n.emissao) == (C2, "2", "2026-09", "2026-09-15")
    assert n.prestador_doc == DONO and n.tomador_doc == CLIENTE and n.tomador_nome == "Cliente Um SA"
    assert n.iss_retido and n.ret_iss == 500.0 and n.aliquota_iss == 5.0
    assert (n.ret_irrf, n.ret_pis, n.ret_cofins, n.ret_csll, n.ret_inss) == (150.0, 65.0, 300.0, 100.0, 0.0)
    assert n.total_retido == 1115.0 and n.valor_liquido == 8885.0 and not n.avisos


def test_iss_nao_retido_e_pis_nao_retido_ficam_zerados():
    n = nfse.ler_nfse(nota_xml(C1, "1", DONO, CLIENTE, "X", 1000.0, iss=50.0, pis=6.5, cofins=30.0, tp_ret_pc="2"))
    assert n.valor_iss == 50.0 and not n.iss_retido and n.ret_iss == 0
    assert n.ret_pis == n.ret_cofins == 0 and n.total_retido == 0


def test_avisa_quando_o_total_retido_nao_bate():
    n = nfse.ler_nfse(nota_xml(C1, "1", DONO, CLIENTE, "X", 1000.0, irrf=15.0, total_ret=20.0))
    assert n.total_retido == 20.0 and n.avisos


def test_evento_nao_e_nota():
    assert nfse.ler_nfse(evento_xml(C4)) is None
    assert nfse.chave_do_evento(evento_xml(C4)) == (C4, "e101101")


def test_busca_por_nsu_ate_acabar_e_continua_do_ultimo(tmp_path):
    portal = PortalFalso(documentos_exemplo())
    cache = nfse.Cache(DONO, tmp_path)
    assert nfse.atualizar_cache(cliente(portal), cache) == 6
    assert [p[0].rsplit("/", 1)[-1] for p in portal.pedidos] == ["0", "2", "4", "6"]
    assert portal.pedidos[0][1] == {"lote": "true"}
    # de novo: começa do último NSU guardado e não traz nada repetido
    portal.pedidos.clear()
    assert nfse.atualizar_cache(cliente(portal), nfse.Cache(DONO, tmp_path)) == 0
    assert portal.pedidos[0][0].endswith("/DFe/6")


def test_espera_e_repete_quando_o_portal_limita(tmp_path):
    esperas = []
    portal = PortalFalso(documentos_exemplo()[:1], falhas=[Resposta(429, {}, cabecalhos={"Retry-After": "3"}),
                                                          Resposta(503, {})])
    c = nfse.ClienteADN(portal, intervalo=0.01, dormir=esperas.append)
    assert nfse.atualizar_cache(c, nfse.Cache(DONO, tmp_path)) == 1
    assert 3.0 in esperas and c.intervalo > 0.01


def test_erro_do_portal_vira_mensagem_clara():
    portal = PortalFalso([], falhas=[Resposta(400, {"Erros": [{"Codigo": "E2214", "Descricao": "Erro Cadeia"}]})])
    with pytest.raises(nfse.ErroADN, match="E2214"):
        cliente(portal).lote(0)
    portal = PortalFalso([], falhas=[Resposta(403, {"Erros": [{"Codigo": "E0001", "Descricao": "Negado"}]})])
    with pytest.raises(nfse.ErroADN, match="não aceitou o certificado"):
        cliente(portal).lote(0)


def test_404_sem_documentos_encerra_a_busca():
    portal = PortalFalso([], falhas=[Resposta(404, {"StatusProcessamento": "NENHUM_DOCUMENTO_LOCALIZADO"})])
    assert list(cliente(portal).documentos(10)) == []


def test_baixa_a_competencia_separa_e_monta_o_relatorio(tmp_path):
    from openpyxl import load_workbook

    portal = PortalFalso(documentos_exemplo())
    r = nfse.baixar_competencia(cliente(portal), DONO, "2026-09", tmp_path / "destino", "Lotivisse", "7",
                                cache_raiz=tmp_path / "cache")
    assert [n.numero for n in r.emitidas] == ["1", "2"]          # a cancelada e a de outubro ficam de fora
    assert [n.numero for n in r.recebidas] == ["77"]
    assert r.canceladas == 1 and r.pdfs == 3
    emitidas = r.pasta / nfse.PASTA_EMITIDAS
    assert r.pasta.parent.name == "7-Lotivisse" and r.pasta.name == "2026-09"
    assert len(list(emitidas.glob("*.xml"))) == 2 and len(list(emitidas.glob("*.pdf"))) == 2
    assert len(list((r.pasta / nfse.PASTA_RECEBIDAS).glob("*.xml"))) == 1

    wb = load_workbook(r.relatorio)
    assert wb.sheetnames == ["Resumo", "Emitidas", "Recebidas"]
    resumo = {linha[0]: linha[1:3] for linha in wb["Resumo"].iter_rows(values_only=True) if linha and linha[0]}
    assert resumo["Quantidade de notas"] == (2, 1)
    assert resumo["ISS retido"] == (500.0, 0.0) and resumo["INSS"] == (0.0, 220.0)
    assert resumo["Total retido"] == (1115.0, 220.0)
    em = list(wb["Emitidas"].iter_rows(values_only=True))
    assert em[0][3] == "CNPJ/CPF do tomador" and em[2][9] == "Sim" and em[-1][0] == "TOTAL"
    assert em[-1][16] == "=SUM(Q2:Q3)"


def test_sem_notas_na_competencia_nao_cria_relatorio(tmp_path):
    r = nfse.baixar_competencia(cliente(PortalFalso([])), DONO, "2026-09", tmp_path / "d", "X",
                                cache_raiz=tmp_path / "c")
    assert r.relatorio is None and not r.emitidas and not (tmp_path / "d").exists()


def _pfx(cn: str, senha: str = "1234") -> bytes:
    from cryptography import x509
    from cryptography.hazmat.primitives import hashes
    from cryptography.hazmat.primitives.asymmetric import rsa
    from cryptography.hazmat.primitives.serialization import BestAvailableEncryption, pkcs12
    from cryptography.x509.oid import NameOID

    chave = rsa.generate_private_key(public_exponent=65537, key_size=2048)
    nome = x509.Name([x509.NameAttribute(NameOID.COMMON_NAME, cn)])
    agora = datetime.now(timezone.utc)
    cert = (x509.CertificateBuilder().subject_name(nome).issuer_name(nome).public_key(chave.public_key())
            .serial_number(1).not_valid_before(agora).not_valid_after(agora + timedelta(days=365))
            .sign(chave, hashes.SHA256()))
    return pkcs12.serialize_key_and_certificates(b"a1", chave, cert, None, BestAvailableEncryption(senha.encode()))


def test_le_cnpj_do_certificado_e_abre_o_tls():
    pfx = _pfx(f"LOTIVISSE LTDA:{DONO}")
    info = nfse.dados_do_certificado(pfx, "1234")
    assert info["cnpj"] == DONO and info["nome"] == "LOTIVISSE LTDA"
    with nfse.contexto_tls(pfx, "1234") as ctx:
        assert ctx is not None
    with pytest.raises(nfse.ErroADN, match="senha"):
        nfse.dados_do_certificado(pfx, "errada")
