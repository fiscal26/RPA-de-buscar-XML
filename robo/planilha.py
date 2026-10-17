"""Planilha de empresas: leitura (modo local) e geração do modelo.

Colunas, as mesmas no portal:
    A  CNPJ                    obrigatório (com ou sem pontuação)
    B  Apelido                 obrigatório (nome curto para achar a pasta)
    C  Senha ATF               obrigatória na primeira vez; no portal, em branco = usa a já cadastrada
    D  Login ATF (opcional)    se o login do ATF não for o próprio CNPJ

No Excel, use a aba "Empresas" a partir da linha 2. Em .csv/.txt: CNPJ;APELIDO;SENHA;LOGIN.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

ABA = "Empresas"
CABECALHO = ["CNPJ", "Apelido", "Senha ATF", "Login ATF (opcional)"]


@dataclass
class Empresa:
    cnpj: str
    apelido: str
    senha: str
    login: str
    linha: int

    @property
    def login_atf(self) -> str:
        return self.login or self.cnpj


@dataclass
class Leitura:
    empresas: list[Empresa] = field(default_factory=list)
    erros: list[str] = field(default_factory=list)
    avisos: list[str] = field(default_factory=list)


def so_digitos(s) -> str:
    return re.sub(r"\D", "", str(s or ""))


def _dv(numeros: str, pesos: list[int]) -> int:
    resto = sum(int(n) * p for n, p in zip(numeros, pesos)) % 11
    return 0 if resto < 2 else 11 - resto


def cnpj_valido(cnpj: str) -> bool:
    if not re.fullmatch(r"\d{14}", cnpj) or len(set(cnpj)) == 1:
        return False
    d1 = _dv(cnpj[:12], [5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2])
    d2 = _dv(cnpj[:13], [6, 5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2])
    return cnpj[-2:] == f"{d1}{d2}"


def normalizar_cnpj(valor) -> tuple[str | None, str | None]:
    """CNPJ só com números, recolocando os zeros à esquerda que o Excel apaga."""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    digitos = so_digitos(valor)
    if not digitos:
        return None, f"'{valor}' não é CNPJ" if str(valor or "").strip() else "CNPJ vazio"
    if len(digitos) > 14:
        return None, f"'{valor}' tem mais de 14 dígitos"
    cnpj = digitos.zfill(14)
    if not cnpj_valido(cnpj):
        return None, f"CNPJ {digitos} inválido (dígito verificador não confere — confira a digitação)"
    return cnpj, None


def formatar_cnpj(cnpj: str) -> str:
    return f"{cnpj[:2]}.{cnpj[2:5]}.{cnpj[5:8]}/{cnpj[8:12]}-{cnpj[12:]}" if len(cnpj) == 14 else cnpj


def _texto(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def ler(caminho: str | Path) -> Leitura:
    caminho = Path(caminho)
    if caminho.suffix.lower() in (".xlsx", ".xlsm"):
        linhas = _linhas_excel(caminho)
    else:
        linhas = _linhas_texto(caminho)
    return interpretar(linhas)


def _linhas_excel(caminho: Path):
    from openpyxl import load_workbook

    wb = load_workbook(caminho, read_only=True, data_only=True)
    ws = wb[ABA] if ABA in wb.sheetnames else wb.worksheets[0]
    for n, valores in enumerate(ws.iter_rows(values_only=True), start=1):
        yield n, list(valores[:4])


def _linhas_texto(caminho: Path):
    bruto = caminho.read_bytes()
    try:
        texto = bruto.decode("utf-8-sig")
    except UnicodeDecodeError:
        texto = bruto.decode("cp1252")  # Bloco de Notas / Excel no Windows
    for n, linha in enumerate(texto.splitlines(), start=1):
        linha = linha.strip()
        if not linha or linha.startswith("#"):
            continue
        sep = next((s for s in (";", "\t", ",") if s in linha), None)
        campos = linha.split(sep) if sep else [linha]
        yield n, [c.strip().strip('"') for c in campos]


def interpretar(linhas) -> Leitura:
    leitura = Leitura()
    vistos: set[str] = set()
    for n, valores in linhas:
        valores = (list(valores) + [None] * 4)[:4]
        cnpj_b, apelido, senha, login = (_texto(v) for v in valores)
        if not any((cnpj_b, apelido, senha, login)):
            continue
        if re.search(r"cnpj", cnpj_b, re.I) and not re.search(r"\d{5}", cnpj_b):
            continue  # cabeçalho
        cnpj, erro = normalizar_cnpj(valores[0])
        if not cnpj:
            leitura.erros.append(f"Linha {n}: {erro}")
            continue
        if not apelido:
            leitura.erros.append(f"Linha {n}: informe o apelido da empresa {formatar_cnpj(cnpj)}")
            continue
        if cnpj in vistos:
            leitura.avisos.append(f"Linha {n}: {formatar_cnpj(cnpj)} repetido; usado só uma vez")
            continue
        vistos.add(cnpj)
        leitura.empresas.append(Empresa(cnpj, apelido[:60], senha, login, n))
    return leitura


def gerar_modelo(destino: str | Path) -> Path:
    """Cria o modelo .xlsx (aba Empresas + aba Instruções). Coluna do CNPJ como texto,
    para o Excel não apagar os zeros à esquerda."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    destino = Path(destino)
    wb = Workbook()
    ws = wb.active
    ws.title = ABA
    ws.append(CABECALHO)
    azul = PatternFill("solid", fgColor="1F4E78")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = azul
        c.alignment = Alignment(vertical="center")
    for col, largura in zip("ABCD", (24, 34, 22, 24)):
        ws.column_dimensions[col].width = largura
    for col in "AD":  # coluna inteira como texto (sem criar linhas vazias)
        ws.column_dimensions[col].number_format = "@"
    ws.freeze_panes = "A2"
    dv = DataValidation(type="textLength", operator="lessThanOrEqual", formula1="60", allow_blank=True,
                        errorTitle="Apelido longo", error="Use até 60 caracteres no apelido.")
    ws.add_data_validation(dv)
    dv.add("B2:B1000")

    inst = wb.create_sheet("Instruções")
    inst.column_dimensions["A"].width = 110
    textos = [
        ("Como preencher", True),
        ("1. Use a aba 'Empresas', uma linha por empresa, a partir da linha 2. Não mude a ordem das colunas.", False),
        ("2. CNPJ: com ou sem pontuação. O portal confere o dígito verificador.", False),
        ("3. Apelido: nome curto da empresa. Vira o nome da pasta onde os XML são salvos.", False),
        ("4. Senha ATF: a senha de acesso ao ATF da SEFAZ-PB. Ela é cifrada ao chegar no portal;", False),
        ("   depois de cadastrada, pode deixar em branco nos próximos envios.", False),
        ("5. Login ATF: preencha só se o login do ATF não for o próprio CNPJ.", False),
        ("   Exemplo de linha:  12.345.678/0001-95 | Padaria Centro | minhaSenha123 | (vazio)", False),
        ("", False),
        ("Segurança", True),
        ("Esta planilha contém senhas. Depois de enviar, apague o arquivo do computador ou guarde em local protegido.", False),
    ]
    for n, (texto, negrito) in enumerate(textos, start=1):
        inst.cell(n, 1, texto).font = Font(bold=negrito, size=12 if negrito else 11)

    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)
    return destino


if __name__ == "__main__":
    import sys

    alvo = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parent.parent / "modelos" / "modelo_empresas_xml.xlsx"
    print(gerar_modelo(alvo))
