"""Arquivos baixados do ATF: guardados como vieram (ZIP continua ZIP) na pasta da empresa.

Pasta final:  <destino>/Arquivos baixados/<ID>-<Nome>/<AAAA-MM>/<NF-e Saida | NF-e Entrada | NFC-e> - R$ <total>/
(faturas, extrato, omissão e certidão: <AAAA-MM>/<Serviço>/, com o PDF como veio do ATF)
(ex.: ID 5 e nome "joao" -> "5-joao"; sem ID na planilha, só o nome). O conteúdo de cada arquivo é
lido só para saber de qual empresa e tipo ele é (IE/CNPJ do emitente ou destinatário dentro das
notas), contar as notas e somar o valor (notas autorizadas, sem as canceladas) que vai no nome da pasta.
Arquivo de empresa fora da rodada fica em "Arquivos baixados/_sem_empresa" e o que não é XML/ZIP em
"_nao_identificados": nada do que foi baixado é descartado.
"""
from __future__ import annotations

import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from xml.etree import ElementTree as ET

from planilha import formatar_ie

# os mesmos nomes das colunas da planilha modelo
TIPOS = {
    "nfe_saida": "NF-e Saida",
    "nfe_entrada": "NF-e Entrada",
    "nfce": "NFC-e",
    "cte_rem": "CT-e Remetente",
    "cte_dest": "CT-e Destinatario",
    "faturas": "Faturas",
    "extrato": "Extrato de Pagamentos",
    "omissao": "Omissão e Inadimplência",
    "certidao": "Certidão de Débitos",
    "cartao": "Conferência Cartão de Crédito",
    "mei": "Dossiê MEI",
}
# Limite anual de receita bruta do MEI (LC 123/2006, art. 18-A, § 1º)
LIMITE_MEI = 81_000.00
# coluna da planilha que vira mais de uma consulta (CT-e: a empresa como remetente e como destinatária)
SUBTIPOS = {"cte": ("cte_rem", "cte_dest")}


def expandir_tipos(tipos: list[str]) -> list[str]:
    return [sub for t in tipos for sub in SUBTIPOS.get(t, (t,))]


# serviços que o ATF responde na hora com PDF (ou ZIP): o arquivo vai direto para a empresa
SERVICOS_DOCUMENTO = ("faturas", "extrato", "omissao", "certidao", "cartao")
PASTA_BAIXADOS = "Arquivos baixados"
_MAX_ARQUIVOS_ZIP = 50_000
_MAX_TAMANHO_XML = 20 * 1024 * 1024


@dataclass
class InfoXml:
    chave: str = ""
    modelo: str = ""          # 55 = NF-e, 65 = NFC-e, 57 = CT-e
    cnpj_emit: str = ""
    cnpj_dest: str = ""
    ie_emit: str = ""
    ie_dest: str = ""
    cnpj_rem: str = ""         # CT-e: remetente
    ie_rem: str = ""
    emissao: str = ""          # aaaa-mm-dd
    evento: bool = False       # cancelamento, carta de correção etc.
    tipo_evento: str = ""      # 110111 = cancelamento
    valor: float = 0.0         # vNF (valor total da nota)
    situacao: str = ""         # cStat da autorização (100 = autorizada)
    partes: list = field(default_factory=list)  # CT-e: (CNPJ, IE) de remetente, destinatário, tomador...


@dataclass
class Resumo:
    salvos: int = 0
    repetidos: int = 0
    conferir: int = 0
    invalidos: int = 0
    pastas: set[str] = field(default_factory=set)

    @property
    def total(self) -> int:
        return self.salvos + self.repetidos

    def somar(self, outro: "Resumo"):
        self.salvos += outro.salvos
        self.repetidos += outro.repetidos
        self.conferir += outro.conferir
        self.invalidos += outro.invalidos
        self.pastas |= outro.pastas


def extrair(dados: bytes, nome: str = "") -> list[tuple[str, bytes]]:
    """Devolve (nome, conteúdo) de cada XML, abrindo ZIP (inclusive ZIP dentro de ZIP) e GZIP."""
    if dados[:2] == b"\x1f\x8b":
        import gzip
        try:
            return extrair(gzip.decompress(dados), Path(nome).stem)
        except Exception:
            return []
    if dados[:4] == b"PK\x03\x04":
        saida: list[tuple[str, bytes]] = []
        with zipfile.ZipFile(io.BytesIO(dados)) as z:
            for i, item in enumerate(z.infolist()):
                if i >= _MAX_ARQUIVOS_ZIP:
                    break
                if item.is_dir() or item.file_size > _MAX_TAMANHO_XML * 5:
                    continue
                conteudo = z.read(item)
                saida.extend(extrair(conteudo, item.filename))
        return saida
    if _parece_xml(dados):
        return [(Path(nome).name or "documento.xml", dados)]
    return []


def _parece_xml(dados: bytes) -> bool:
    inicio = dados[:200].lstrip(b"\xef\xbb\xbf \r\n\t")
    return inicio.startswith(b"<")


def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def ler_xml(dados: bytes) -> InfoXml | None:
    """Lê os dados que importam de uma NF-e/NFC-e (nfeProc, NFe) ou evento (procEventoNFe)."""
    if len(dados) > _MAX_TAMANHO_XML:
        return None
    try:
        raiz = ET.fromstring(dados)
    except ET.ParseError:
        return None
    info = InfoXml()
    nomes = {_local(e.tag): e for e in raiz.iter()}  # primeira ocorrência basta para estes campos
    raiz_nome = _local(raiz.tag)

    if "infNFe" in nomes:
        inf = nomes["infNFe"]
        info.chave = re.sub(r"\D", "", inf.get("Id", ""))
        ide = _filho(inf, "ide")
        info.modelo = _texto(ide, "mod")
        info.emissao = (_texto(ide, "dhEmi") or _texto(ide, "dEmi"))[:10]
        emit, dest = _filho(inf, "emit"), _filho(inf, "dest")
        info.cnpj_emit = _texto(emit, "CNPJ") or _texto(emit, "CPF")
        info.cnpj_dest = _texto(dest, "CNPJ") or _texto(dest, "CPF")
        info.ie_emit = re.sub(r"\D", "", _texto(emit, "IE"))
        info.ie_dest = re.sub(r"\D", "", _texto(dest, "IE"))
        try:
            info.valor = float(_texto(_filho(_filho(inf, "total"), "ICMSTot"), "vNF") or 0)
        except ValueError:
            info.valor = 0.0
        prot = nomes.get("infProt")
        info.situacao = _texto(prot, "cStat") if prot is not None else ""
    elif "infCte" in nomes:  # CT-e (cteProc): o valor é o vTPrest
        inf = nomes["infCte"]
        info.chave = re.sub(r"\D", "", inf.get("Id", ""))
        ide = _filho(inf, "ide")
        info.modelo = _texto(ide, "mod") or "57"
        info.emissao = (_texto(ide, "dhEmi") or _texto(ide, "dEmi"))[:10]
        emit, dest = _filho(inf, "emit"), _filho(inf, "dest")
        info.cnpj_emit = _texto(emit, "CNPJ") or _texto(emit, "CPF")
        info.cnpj_dest = _texto(dest, "CNPJ") or _texto(dest, "CPF")
        info.ie_emit = re.sub(r"\D", "", _texto(emit, "IE"))
        info.ie_dest = re.sub(r"\D", "", _texto(dest, "IE"))
        rem = _filho(inf, "rem")
        info.cnpj_rem = (_texto(rem, "CNPJ") or _texto(rem, "CPF")) if rem is not None else ""
        info.ie_rem = re.sub(r"\D", "", _texto(rem, "IE")) if rem is not None else ""
        for papel in ("rem", "dest", "exped", "receb"):
            e = _filho(inf, papel)
            if e is not None:
                info.partes.append((_texto(e, "CNPJ") or _texto(e, "CPF"), re.sub(r"\D", "", _texto(e, "IE"))))
        toma4 = _filho(ide, "toma4")
        if toma4 is not None:
            info.partes.append((_texto(toma4, "CNPJ") or _texto(toma4, "CPF"), re.sub(r"\D", "", _texto(toma4, "IE"))))
        try:
            info.valor = float(_texto(_filho(inf, "vPrest"), "vTPrest") or 0)
        except ValueError:
            info.valor = 0.0
        prot = nomes.get("infProt")
        info.situacao = _texto(prot, "cStat") if prot is not None else ""
    elif "infEvento" in nomes or raiz_nome in ("procEventoNFe", "procEventoCTe", "evento", "eventoCTe"):
        inf = nomes.get("infEvento")
        info.evento = True
        info.chave = _texto(inf, "chNFe") or _texto(inf, "chCTe")
        info.tipo_evento = _texto(inf, "tpEvento")
        info.cnpj_emit = _texto(inf, "CNPJ") or _texto(inf, "CPF")
        info.emissao = _texto(inf, "dhEvento")[:10]
        info.modelo = info.chave[20:22] if len(info.chave) == 44 else ""
    else:
        return None
    if not re.fullmatch(r"\d{44}", info.chave or ""):
        return None
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", info.emissao):
        # sem data legível: usa ano e mês da própria chave (posições 3 a 6 = AAMM)
        info.emissao = f"20{info.chave[2:4]}-{info.chave[4:6]}-01"
    return info


def _filho(elem, nome):
    if elem is None:
        return None
    for e in elem:
        if _local(e.tag) == nome:
            return e
    return None


def _texto(elem, nome) -> str:
    e = _filho(elem, nome)
    return (e.text or "").strip() if e is not None and e.text else ""


def nome_de_pasta(texto: str) -> str:
    texto = unicodedata.normalize("NFKC", texto).strip()
    texto = re.sub(r'[<>:"/\\|?*\x00-\x1f]+', " ", texto)  # caracteres proibidos no Windows
    return re.sub(r"\s+", " ", texto).strip(" .")[:60] or "Empresa"


def pasta_empresa(destino: str | Path, nome: str, ie: str, codigo: str = "") -> Path:
    """<destino>/Arquivos baixados/<ID>-<Nome> (ex.: "5-joao"); sem ID, só o nome."""
    base = Path(destino) / PASTA_BAIXADOS
    return base / nome_de_pasta(f"{codigo}-{nome}" if codigo else nome)


def _mesmo_ie(a: str, b: str) -> bool:
    return bool(a) and bool(b) and a.lstrip("0") == b.lstrip("0")


def pertence(info: InfoXml, ie: str, cnpj: str, tipo: str) -> bool:
    """O XML é da empresa? Saída/NFC-e: ela é a emitente. Entrada: ela é a destinatária.
    Confere pela IE e, se houver, pelo CNPJ. Evento só traz o autor e a chave (CNPJ do emitente
    nas posições 7 a 20)."""
    emitente_chave = info.chave[6:20]
    if tipo != "nfe_entrada":
        if info.evento:
            return not cnpj or cnpj in (info.cnpj_emit, emitente_chave)
        return _mesmo_ie(info.ie_emit, ie) or (bool(cnpj) and cnpj in (info.cnpj_emit, emitente_chave))
    if info.evento:
        return not cnpj or info.cnpj_emit == cnpj or emitente_chave != cnpj
    return _mesmo_ie(info.ie_dest, ie) or (bool(cnpj) and info.cnpj_dest == cnpj)


def salvar(arquivos: list[tuple[str, bytes]], destino: str | Path, nome: str, ie: str, cnpj: str, tipo: str,
           codigo: str = "") -> Resumo:
    resumo = Resumo()
    base = pasta_empresa(destino, nome, ie, codigo)
    nome_tipo = TIPOS.get(tipo, tipo)
    for nome, dados in arquivos:
        info = ler_xml(dados)
        if not info:
            resumo.invalidos += 1
            continue
        mes = info.emissao[:7]
        pasta = base / mes / nome_tipo
        conferir = not pertence(info, ie, cnpj, tipo)
        if conferir:
            pasta = base / "_conferir"
        if info.evento:
            pasta = pasta / "Eventos"
        arquivo = pasta / (f"{info.chave}-evento-{_sufixo(nome)}.xml" if info.evento else f"{info.chave}.xml")
        resumo.pastas.add(str(base))
        repetido = arquivo.exists() and arquivo.read_bytes() == dados
        if not repetido:
            pasta.mkdir(parents=True, exist_ok=True)
            arquivo.write_bytes(dados)
        if conferir:
            resumo.conferir += 1  # conta sempre: o portal avisa que há XML de outro CNPJ
        elif repetido:
            resumo.repetidos += 1
        else:
            resumo.salvos += 1
    return resumo


def _sufixo(nome: str) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "", Path(nome).stem)[-20:] or "1"


# ---------------------------------------------------------------------- distribuição por conteúdo
# O ATF entrega os XML como anexos em "Minhas Mensagens", todas com o mesmo assunto e sem ordem
# garantida. Por isso cada XML é encaminhado pelo que está DENTRO dele (IE/CNPJ de emitente e
# destinatário, modelo 55/65), e não pela mensagem de onde veio.

@dataclass
class Alvo:
    """Empresa consultada nesta rodada e os XML que ela pediu."""
    ie: str
    cnpj: str
    tipos: list[str]


def _destinos(info: InfoXml, alvos: list[Alvo]) -> list[tuple[str, str]]:
    emit_chave = info.chave[6:20]
    saida: list[tuple[str, str]] = []
    for a in alvos:
        emitente = _mesmo_ie(info.ie_emit, a.ie) or (bool(a.cnpj) and a.cnpj in (info.cnpj_emit, emit_chave))
        if info.modelo == "57":  # CT-e: a empresa como remetente e/ou como destinatária
            if "cte_rem" in a.tipos and (_mesmo_ie(info.ie_rem, a.ie) or (bool(a.cnpj) and info.cnpj_rem == a.cnpj)):
                saida.append((a.ie, "cte_rem"))
            if "cte_dest" in a.tipos and (_mesmo_ie(info.ie_dest, a.ie) or (bool(a.cnpj) and info.cnpj_dest == a.cnpj)):
                saida.append((a.ie, "cte_dest"))
            continue
        if info.modelo == "65":
            if "nfce" in a.tipos and emitente:
                saida.append((a.ie, "nfce"))
            continue
        if "nfe_saida" in a.tipos and emitente:
            saida.append((a.ie, "nfe_saida"))
        if "nfe_entrada" in a.tipos and (_mesmo_ie(info.ie_dest, a.ie) or (bool(a.cnpj) and info.cnpj_dest == a.cnpj)):
            saida.append((a.ie, "nfe_entrada"))
    return saida


def distribuir(arquivos: list[tuple[str, bytes]], alvos: list[Alvo]):
    """Separa os XML recebidos por (IE, tipo). Uma nota entre duas empresas da rodada vai para as
    duas (saída de uma, entrada da outra). Eventos seguem a nota da mesma chave. Devolve
    ({(ie, tipo): [(nome, dados)]}, sobras): sobras são XML de nenhuma empresa da rodada."""
    por_destino: dict[tuple[str, str], list[tuple[str, bytes]]] = {}
    sobras: list[tuple[str, bytes]] = []
    destinos_da_chave: dict[str, list[tuple[str, str]]] = {}
    eventos = []
    for nome, dados in arquivos:
        info = ler_xml(dados)
        if not info:
            continue
        if info.evento:
            eventos.append((nome, dados, info))
            continue
        destinos = _destinos(info, alvos)
        da_chave = destinos_da_chave.setdefault(info.chave, [])
        da_chave.extend(d for d in destinos if d not in da_chave)
        for d in destinos:
            por_destino.setdefault(d, []).append((nome, dados))
        if not destinos:
            sobras.append((nome, dados))
    for nome, dados, info in eventos:
        destinos = destinos_da_chave.get(info.chave)
        if destinos is None and info.modelo == "57":
            destinos = []  # evento de CT-e sem o CT-e: o CNPJ da chave é o da transportadora
        elif destinos is None:  # nota não veio nesta rodada: pelo CNPJ do emitente na chave
            destinos = [(a.ie, "nfce" if info.modelo == "65" else "nfe_saida") for a in alvos
                        if a.cnpj and a.cnpj == info.chave[6:20]
                        and ("nfce" if info.modelo == "65" else "nfe_saida") in a.tipos]
        for d in destinos:
            por_destino.setdefault(d, []).append((nome, dados))
        if not destinos:
            sobras.append((nome, dados))
    return por_destino, sobras


def salvar_sobras(arquivos: list[tuple[str, bytes]], destino: str | Path) -> int:
    """XML que não é de nenhuma empresa da rodada (ex.: consulta feita à mão no mesmo login)."""
    n = 0
    for nome, dados in arquivos:
        info = ler_xml(dados)
        if not info:
            continue
        pasta = Path(destino) / PASTA_BAIXADOS / "_sem_empresa" / info.emissao[:7]
        arquivo = pasta / (f"{info.chave}-evento-{_sufixo(nome)}.xml" if info.evento else f"{info.chave}.xml")
        if not arquivo.exists():
            pasta.mkdir(parents=True, exist_ok=True)
            arquivo.write_bytes(dados)
            n += 1
    return n


def salvar_nao_identificados(arquivos: list[tuple[str, bytes]], destino: str | Path,
                              subpasta: str = "_nao_identificados") -> list[Path]:
    """Arquivo baixado que não é XML nem ZIP de XML: guarda como veio, para conferir (nada se perde)."""
    import hashlib

    pasta = Path(destino) / PASTA_BAIXADOS / subpasta
    salvos = []
    for nome, dados in arquivos:
        base = Path(nome_de_pasta(Path(nome).name) or "arquivo")
        arquivo = pasta / f"{base.stem}_{hashlib.sha1(dados).hexdigest()[:8]}{base.suffix}"
        if not arquivo.exists():
            pasta.mkdir(parents=True, exist_ok=True)
            arquivo.write_bytes(dados)
        salvos.append(arquivo)
    return salvos


def tipo_de_arquivo(dados: bytes) -> str:
    if dados[:4] == b"PK\x03\x04":
        return "ZIP"
    if dados[:2] == b"\x1f\x8b":
        return "GZIP"
    inicio = dados[:300].lstrip(b"\xef\xbb\xbf \r\n\t").lower()
    if inicio.startswith(b"<?xml") or inicio.startswith(b"<nfeproc") or inicio.startswith(b"<nfe"):
        return "XML"
    if inicio.startswith(b"<!doctype html") or inicio.startswith(b"<html"):
        return "HTML"
    if dados[:4] == b"%PDF":
        return "PDF"
    return "outro"


# ---------------------------------------------------------------------- arquivos como vieram do ATF
# O ZIP baixado é guardado do jeito que chegou (sem descompactar). O conteúdo só é lido para saber
# de qual empresa e tipo ele é, contar as notas e somar o valor que aparece no nome da pasta.

def distribuir_arquivos(brutos: list[tuple[str, bytes]], alvos: list["Alvo"]):
    """Cada arquivo baixado (ZIP ou XML) vai inteiro para a empresa/tipo da maioria das suas notas.
    Devolve ({(ie, tipo): [(nome, dados)]}, sem_empresa, nao_identificados)."""
    por_destino: dict[tuple[str, str], list[tuple[str, bytes]]] = {}
    sem_empresa, nao_identificados, empatados = [], [], []
    for nome, dados in brutos:
        xmls = extrair(dados, nome)
        if not xmls:
            nao_identificados.append((nome, dados))
            continue
        por, _ = distribuir(xmls, alvos)
        if not por:
            sem_empresa.append((nome, dados))
            continue
        maior = max(len(v) for v in por.values())
        candidatos = [k for k, v in por.items() if len(v) == maior]
        if len(candidatos) == 1:
            por_destino.setdefault(candidatos[0], []).append((nome, dados))
        else:  # ex.: ZIP só com notas entre duas empresas da rodada (saída de uma = entrada da outra)
            empatados.append((nome, dados, candidatos))
    # empate: cada consulta gera o seu arquivo, então vai para quem ainda não recebeu nenhum
    for nome, dados, candidatos in empatados:
        livre = next((k for k in candidatos if k not in por_destino), candidatos[0])
        por_destino.setdefault(livre, []).append((nome, dados))
    return por_destino, sem_empresa, nao_identificados


def valor_total(xmls: list[tuple[str, bytes]]) -> tuple[int, float]:
    """(quantidade de notas, soma do vNF) das notas autorizadas e não canceladas."""
    notas: dict[str, float] = {}
    canceladas: set[str] = set()
    for _, dados in xmls:
        info = ler_xml(dados)
        if not info:
            continue
        if info.evento:
            if info.tipo_evento in ("110111", "110112"):  # cancelamento (e por substituição)
                canceladas.add(info.chave)
        elif info.situacao in ("", "100", "150"):
            notas[info.chave] = info.valor
    validas = {k: v for k, v in notas.items() if k not in canceladas}
    return len(validas), round(sum(validas.values()), 2)


def reais(valor: float) -> str:
    return "R$ " + f"{valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _pasta_do_tipo(pasta_mes: Path, nome_tipo: str) -> Path:
    """A pasta do tipo pode já existir com o valor no nome ("NF-e Saida - R$ 1.234,56")."""
    if pasta_mes.is_dir():
        for d in sorted(pasta_mes.iterdir()):
            if d.is_dir() and (d.name == nome_tipo or d.name.startswith(nome_tipo + " - R$")):
                return d
    return pasta_mes / nome_tipo


def salvar_originais(brutos: list[tuple[str, bytes]], destino: str | Path, nome: str, ie: str, tipo: str,
                     competencia: str, codigo: str = "") -> dict:
    """Guarda os arquivos baixados como vieram (ZIP continua ZIP) em
    Arquivos baixados/<ID>-<Nome>/<AAAA-MM>/<tipo - R$ total>/ e devolve
    {novos, repetidos, notas, valor, pasta}. Rodar de novo não duplica (compara o conteúdo)."""
    import hashlib

    nome_tipo = TIPOS.get(tipo, tipo)
    pasta_mes = pasta_empresa(destino, nome, ie, codigo) / competencia
    pasta = _pasta_do_tipo(pasta_mes, nome_tipo)
    existentes = {hashlib.sha1(a.read_bytes()).hexdigest() for a in pasta.iterdir() if a.is_file()} if pasta.is_dir() else set()
    novos = repetidos = 0
    for nome_arq, dados in brutos:
        assinatura = hashlib.sha1(dados).hexdigest()
        if assinatura in existentes:
            repetidos += 1
            continue
        existentes.add(assinatura)
        ext = ".zip" if dados[:4] == b"PK\x03\x04" else ".gz" if dados[:2] == b"\x1f\x8b" else ".xml"
        pasta.mkdir(parents=True, exist_ok=True)
        base = f"{nome_tipo} {competencia}"
        arquivo, n = pasta / f"{base}{ext}", 1
        while arquivo.exists():
            n += 1
            arquivo = pasta / f"{base} ({n}){ext}"
        arquivo.write_bytes(dados)
        novos += 1
    if not pasta.is_dir():
        return {"novos": 0, "repetidos": repetidos, "notas": 0, "valor": 0.0, "pasta": ""}
    # valor de tudo o que está na pasta (inclusive de rodadas anteriores) e nome da pasta com o total
    todos = [x for a in sorted(pasta.iterdir()) if a.is_file() for x in extrair(a.read_bytes(), a.name)]
    notas, valor = valor_total(todos)
    novo_nome = f"{nome_tipo} - {reais(valor)}" if notas else nome_tipo
    if pasta.name != novo_nome and not (pasta_mes / novo_nome).exists():
        pasta = pasta.rename(pasta_mes / novo_nome)
    return {"novos": novos, "repetidos": repetidos, "notas": notas, "valor": valor, "pasta": str(pasta)}


def extensao(dados: bytes) -> str:
    if dados[:4] == b"PK\x03\x04":
        return ".zip"
    if dados[:1024].lstrip()[:5] == b"%PDF-":
        return ".pdf"
    if dados[:2] == b"\x1f\x8b":
        return ".gz"
    if dados[:4] == b"\x89PNG":
        return ".png"
    return {"XML": ".xml", "HTML": ".html"}.get(tipo_de_arquivo(dados), ".bin")


def salvar_documentos(brutos: list[tuple[str, bytes]], destino: str | Path, nome: str, ie: str, tipo: str,
                      competencia: str, codigo: str = "") -> dict:
    """Faturas, extrato, omissão e certidão: guarda o PDF (ou ZIP) como veio em
    Arquivos baixados/<ID>-<Nome>/<AAAA-MM>/<Serviço>/<Serviço> <AAAA-MM>.pdf. Rodar de novo guarda o
    novo ao lado ("(2)"), sem repetir arquivo igual. Devolve {novos, repetidos, pasta}."""
    import hashlib

    nome_tipo = TIPOS.get(tipo, tipo)
    pasta = pasta_empresa(destino, nome, ie, codigo) / competencia / nome_tipo
    existentes = {hashlib.sha1(a.read_bytes()).hexdigest() for a in pasta.iterdir() if a.is_file()} if pasta.is_dir() else set()
    novos = repetidos = 0
    for _, dados in brutos:
        assinatura = hashlib.sha1(dados).hexdigest()
        if assinatura in existentes:
            repetidos += 1
            continue
        existentes.add(assinatura)
        ext = extensao(dados)
        pasta.mkdir(parents=True, exist_ok=True)
        base = f"{nome_tipo} {competencia}"
        arquivo, n = pasta / f"{base}{ext}", 1
        while arquivo.exists():
            n += 1
            arquivo = pasta / f"{base} ({n}){ext}"
        arquivo.write_bytes(dados)
        novos += 1
    return {"novos": novos, "repetidos": repetidos, "pasta": str(pasta) if pasta.is_dir() else ""}


def _arquivo_mei(destino, nome, ie, codigo, ano) -> Path:
    return pasta_empresa(destino, nome, ie, codigo) / "Dossiê MEI" / f"Faturamento MEI {ano}.json"


def periodos_mei(destino: str | Path, nome: str, ie: str, codigo: str, inicio, refazer_ano: bool = False) -> list:
    """Meses do ano ANTES da competência que o Dossiê MEI ainda não tem (de janeiro em diante): o cliente
    que chega sem acompanhamento já sai com o ano todo. refazer_ano=True busca todos de novo.
    Devolve [(1º dia, último dia)]."""
    import calendar
    import json
    from datetime import date

    try:
        feitos = json.loads(_arquivo_mei(destino, nome, ie, codigo, inicio.year).read_text(encoding="utf-8"))["meses"]
    except (OSError, ValueError, KeyError):
        feitos = {}
    saida = []
    for mes in range(1, inicio.month):
        if refazer_ano or f"{inicio.year}-{mes:02d}" not in feitos:
            saida.append((date(inicio.year, mes, 1), date(inicio.year, mes, calendar.monthrange(inicio.year, mes)[1])))
    return saida


def limite_mei(ano: int, inicio_atividade: str | None) -> float:
    """R$ 81.000,00 no ano; no ano em que a empresa começou, R$ 6.750,00 por mês de atividade, contando o
    mês de início (LC 123/2006, art. 18-A, §§ 1º e 2º)."""
    if inicio_atividade and int(inicio_atividade[:4]) == ano:
        return round(LIMITE_MEI / 12 * (12 - int(inicio_atividade[5:7]) + 1), 2)
    return LIMITE_MEI


def salvar_faturamento_mei(destino: str | Path, nome: str, ie: str, competencia: str, dados: dict,
                           codigo: str = "") -> dict:
    """Dossiê MEI: guarda a tabela "Totais docs fiscais eletrônicos" do mês em
    Arquivos baixados/<ID>-<Nome>/Dossiê MEI/Faturamento MEI <ano>.json (um mês por chave, base para o
    painel futuro) e refaz a planilha "Faturamento MEI <ano>.xlsx". NOTAS e CARTÃO são acompanhados
    separados (o cartão pode estar ou não dentro das notas): cada um com o seu acumulado e a sua % do
    limite. Abas: o ano mês a mês ("sem movimento" quando tudo é zero), todas as linhas por tipo de
    documento (inclusive as zeradas) e a tabela como veio do ATF."""
    import json
    from datetime import datetime

    ano = competencia[:4]
    arq_json = _arquivo_mei(destino, nome, ie, codigo, ano)
    pasta = arq_json.parent
    pasta.mkdir(parents=True, exist_ok=True)
    try:
        historico = json.loads(arq_json.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        historico = {}
    empresa = historico.setdefault("empresa", {"ie": ie, "nome": nome, "codigo": codigo})
    if dados.get("inicio_atividade"):
        empresa["inicio_atividade"] = dados["inicio_atividade"]
    limite = limite_mei(int(ano), empresa.get("inicio_atividade"))
    empresa["limite_ano"] = limite
    meses = historico.setdefault("meses", {})
    meses[competencia] = {**{k: dados.get(k) for k in ("saidas", "entradas", "qtd_saidas", "qtd_entradas", "cartao",
                                                       "sem_movimento", "simei")},
                          "itens": dados.get("itens", []), "linhas": dados.get("linhas", []),
                          "consultado_em": datetime.now().isoformat(timespec="seconds")}
    arq_json.write_text(json.dumps(historico, ensure_ascii=False, indent=1, sort_keys=True), encoding="utf-8")

    from openpyxl import Workbook
    from openpyxl.styles import Font

    def negrito(ws):
        for c in ws[1]:
            c.font = Font(bold=True)

    wb = Workbook()
    ws = wb.active
    ws.title = f"Faturamento {ano}"
    ws.append(["Mês", "Movimento",
               "Notas: saídas", "Notas: qtd. saídas", "Notas: entradas", "Notas: qtd. entradas",
               "Notas: saídas acumuladas", f"Notas: % do limite ({reais(limite)})",
               "Cartão: vendas", "Cartão: qtd.", "Cartão: acumulado", f"Cartão: % do limite ({reais(limite)})",
               "Consultado em"])
    notas_acum = cartao_acum = 0.0
    for mes in sorted(meses):
        m = meses[mes]
        cartao = m.get("cartao") or {}
        notas_acum += m.get("saidas") or 0.0
        cartao_acum += cartao.get("valor") or 0.0
        ws.append([mes, "sem movimento" if m.get("sem_movimento") else "com movimento",
                   m.get("saidas"), m.get("qtd_saidas"), m.get("entradas"), m.get("qtd_entradas"),
                   round(notas_acum, 2), round(notas_acum / limite, 4),
                   cartao.get("valor"), cartao.get("qtd"), round(cartao_acum, 2), round(cartao_acum / limite, 4),
                   m.get("consultado_em")])
    negrito(ws)
    for linha in ws.iter_rows(min_row=2):
        for i in (2, 4, 6, 8, 10):
            linha[i].number_format = '#,##0.00'
        for i in (7, 11):
            linha[i].number_format = "0.00%"
    for col, w in zip("ABCDEFGHIJKLM", (10, 16, 15, 12, 15, 13, 18, 22, 15, 11, 17, 22, 20)):
        ws.column_dimensions[col].width = w

    docs = wb.create_sheet("Por documento")
    docs.append(["Mês", "Tipo de documento", "Qtd. entradas", "Entradas", "Qtd. saídas", "Saídas"])
    for mes in sorted(meses):
        for i in meses[mes].get("itens", []):
            docs.append([mes, i.get("tipo"), i.get("qtd_entradas"), i.get("entradas"), i.get("qtd_saidas"), i.get("saidas")])
        c = meses[mes].get("cartao")
        if c:
            docs.append([mes, "Cartão (à parte)", None, None, c.get("qtd"), c.get("valor")])
    negrito(docs)
    docs.column_dimensions["B"].width = 22

    bruto = wb.create_sheet("Tabela do ATF")
    for mes in sorted(meses):
        for celulas in meses[mes].get("linhas", []):
            bruto.append([mes] + list(celulas))
    arq_xlsx = pasta / f"Faturamento MEI {ano}.xlsx"
    try:
        wb.save(arq_xlsx)
    except PermissionError:  # planilha aberta no Excel: o JSON já foi salvo; a planilha sai na próxima vez
        pass
    return {"pasta": str(pasta), "saidas": dados.get("saidas"), "entradas": dados.get("entradas"),
            "cartao": (dados.get("cartao") or {}).get("valor"), "sem_movimento": dados.get("sem_movimento"),
            "limite": limite, "meses": sorted(meses),
            "notas_acumulado": round(notas_acum, 2), "notas_percentual": round(notas_acum / limite * 100, 1),
            "cartao_acumulado": round(cartao_acum, 2), "cartao_percentual": round(cartao_acum / limite * 100, 1),
            # o ano inteiro (sem a tabela bruta) para o painel de Monitoramento MEI
            "inicio_atividade": empresa.get("inicio_atividade"),
            "historico": {mes: {k: v for k, v in m.items() if k != "linhas"} for mes, m in sorted(meses.items())}}


def historicos_mei(destino: str | Path) -> list[dict]:
    """Todos os "Faturamento MEI <ano>.json" já salvos neste PC (para mandar ao painel o que veio antes):
    [{"ie", "nome", "inicio_atividade", "meses": {aaaa-mm: {...sem a tabela bruta}}}]."""
    import json

    saida = []
    for arq in sorted((Path(destino) / PASTA_BAIXADOS).glob("*/Dossiê MEI/Faturamento MEI *.json")):
        try:
            h = json.loads(arq.read_text(encoding="utf-8"))
            emp = h["empresa"]
        except (OSError, ValueError, KeyError):
            continue
        meses = {mes: {k: v for k, v in m.items() if k != "linhas"} for mes, m in sorted(h.get("meses", {}).items())}
        if emp.get("ie") and meses:
            saida.append({"ie": emp["ie"], "nome": emp.get("nome", ""), "inicio_atividade": emp.get("inicio_atividade"),
                          "meses": meses})
    return saida


def salvar_sem_empresa(brutos: list[tuple[str, bytes]], destino: str | Path) -> int:
    """Arquivo baixado que não é de nenhuma empresa da rodada: guardado como veio."""
    return len(salvar_nao_identificados(brutos, destino, "_sem_empresa"))
